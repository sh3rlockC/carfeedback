from __future__ import annotations

import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models import ConfirmedVehicleSeries, SeriesConflict, SeriesImportBatch
from app.services.confirmed_vehicle_series import PLATFORMS, query_key
from app.services.series_admin import SeriesMutation, create_series_record, now

IMPORT_SUMMARY_KEYS = ("new", "duplicate", "conflict", "invalid")
REQUIRED_FIELDS = ("query", "platform", "series_id")
OPTIONAL_FIELDS = ("url", "title", "source")
SUPPORTED_EXCEL_SUFFIXES = {".xlsx", ".xlsm"}


@dataclass(frozen=True)
class ImportRowPreview:
    row_number: int
    query: str | None
    platform: str | None
    series_id: str | None
    status: str
    query_key: str | None = None
    url: str | None = None
    title: str | None = None
    source: str | None = None
    error: str | None = None
    existing_value_json: dict[str, Any] | None = None
    incoming_value_json: dict[str, Any] | None = None


@dataclass(frozen=True)
class ImportPreview:
    filename: str
    summary: dict[str, int]
    rows: list[ImportRowPreview]


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_row(raw: dict[str, Any], row_number: int) -> dict[str, str]:
    row = {"_row_number": str(row_number)}
    for field in (*REQUIRED_FIELDS, *OPTIONAL_FIELDS):
        row[field] = _cell_to_text(raw.get(field))
    return row


def _read_csv_rows(content: bytes) -> list[dict[str, str]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(text))
    if reader.fieldnames is None:
        return []
    return [_normalize_row(raw, index) for index, raw in enumerate(reader, start=2)]


def _read_excel_rows(content: bytes) -> list[dict[str, str]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration:
        return []

    headers = [_cell_to_text(value) for value in header_values]
    parsed_rows: list[dict[str, str]] = []
    for index, values in enumerate(rows, start=2):
        raw = {header: value for header, value in zip(headers, values, strict=False) if header}
        parsed_rows.append(_normalize_row(raw, index))
    return parsed_rows


def _read_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return _read_csv_rows(content)
    if suffix in SUPPORTED_EXCEL_SUFFIXES:
        return _read_excel_rows(content)
    raise ValueError("series import supports .csv, .xlsx, and .xlsm files")


def _snapshot(record: ConfirmedVehicleSeries) -> dict[str, Any]:
    return {
        "id": record.id,
        "query_key": record.query_key,
        "query": record.query,
        "platform": record.platform,
        "series_id": record.series_id,
        "status": record.status,
        "url": record.url,
        "title": record.title,
        "source": record.source,
        "import_batch_id": record.import_batch_id,
    }


def _incoming(row: dict[str, str], key: str) -> dict[str, Any]:
    return {
        "query_key": key,
        "query": row["query"],
        "platform": row["platform"],
        "series_id": row["series_id"],
        "url": row["url"] or None,
        "title": row["title"] or None,
        "source": row["source"] or None,
    }


def _active_record(db: Session, *, key: str, platform: str) -> ConfirmedVehicleSeries | None:
    return (
        db.query(ConfirmedVehicleSeries)
        .filter(
            ConfirmedVehicleSeries.query_key == key,
            ConfirmedVehicleSeries.platform == platform,
            ConfirmedVehicleSeries.status == "active",
        )
        .one_or_none()
    )


def _invalid_reason(row: dict[str, str]) -> str | None:
    missing = [field for field in REQUIRED_FIELDS if not row[field]]
    if missing:
        return f"missing required field: {', '.join(missing)}"
    if row["platform"] not in PLATFORMS:
        return "unsupported platform"
    return None


def preview_series_import(db: Session, *, filename: str, content: bytes) -> ImportPreview:
    summary = {key: 0 for key in IMPORT_SUMMARY_KEYS}
    previews: list[ImportRowPreview] = []

    for row in _read_rows(filename, content):
        row_number = int(row["_row_number"])
        error = _invalid_reason(row)
        key = query_key(row["query"]) if row["query"] else None
        if error is not None:
            summary["invalid"] += 1
            previews.append(
                ImportRowPreview(
                    row_number=row_number,
                    query=row["query"] or None,
                    platform=row["platform"] or None,
                    series_id=row["series_id"] or None,
                    query_key=key,
                    url=row["url"] or None,
                    title=row["title"] or None,
                    source=row["source"] or None,
                    status="invalid",
                    error=error,
                )
            )
            continue

        assert key is not None
        existing = _active_record(db, key=key, platform=row["platform"])
        incoming = _incoming(row, key)
        if existing is None:
            status = "new"
            existing_value = None
        elif existing.series_id == row["series_id"]:
            status = "duplicate"
            existing_value = _snapshot(existing)
        else:
            status = "conflict"
            existing_value = _snapshot(existing)

        summary[status] += 1
        previews.append(
            ImportRowPreview(
                row_number=row_number,
                query=row["query"],
                platform=row["platform"],
                series_id=row["series_id"],
                query_key=key,
                url=row["url"] or None,
                title=row["title"] or None,
                source=row["source"] or None,
                status=status,
                existing_value_json=existing_value,
                incoming_value_json=incoming,
            )
        )

    return ImportPreview(filename=filename, summary=summary, rows=previews)


def commit_series_import(db: Session, *, filename: str, content: bytes, operator: str) -> ImportPreview:
    preview = preview_series_import(db, filename=filename, content=content)
    batch = SeriesImportBatch(source="import", filename=filename, operator=operator, summary_json=preview.summary)
    db.add(batch)
    db.flush()

    for row in preview.rows:
        if row.status == "new":
            record = create_series_record(
                db,
                SeriesMutation(
                    query=row.query or "",
                    platform=row.platform or "",
                    series_id=row.series_id or "",
                    url=row.url,
                    title=row.title,
                    source=row.source,
                    operator=operator,
                    reason=f"import batch {batch.id}",
                ),
            )
            record.import_batch_id = batch.id
            db.flush()
        elif row.status == "conflict":
            db.add(
                SeriesConflict(
                    conflict_type="series_import",
                    query_key=row.query_key or "",
                    query=row.query or "",
                    platform=row.platform,
                    existing_value_json=row.existing_value_json or {},
                    incoming_value_json=row.incoming_value_json or {},
                    status="open",
                    import_batch_id=batch.id,
                    created_at=now(),
                )
            )

    db.commit()
    return preview


def export_series_excel(db: Session) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "confirmed_vehicle_series"
    worksheet.append(["query", "platform", "series_id", "url", "title", "source", "status", "updated_at"])

    records = (
        db.query(ConfirmedVehicleSeries)
        .order_by(ConfirmedVehicleSeries.query.asc(), ConfirmedVehicleSeries.platform.asc())
        .all()
    )
    for record in records:
        worksheet.append(
            [
                record.query,
                record.platform,
                record.series_id,
                record.url,
                record.title,
                record.source,
                record.status,
                record.updated_at.isoformat() if record.updated_at else None,
            ]
        )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
