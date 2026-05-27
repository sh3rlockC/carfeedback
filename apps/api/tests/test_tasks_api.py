from __future__ import annotations

import os
from pathlib import Path
import zipfile
from urllib.parse import parse_qs, urlparse
import sys

from fastapi.testclient import TestClient
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

os.environ["DATABASE_URL"] = "sqlite+pysqlite:///:memory:"

from app.config import Settings
from app.db import get_session_local, reset_engine_cache
from app.main import create_app
from app.models import Task, TaskArtifact, TaskEvent
from app.services.passphrase import hash_passphrase
from app.services.task_tokens import hash_task_token
from app.services.task_workflow_client import get_task_workflow_client


class FakeTaskWorkflowClient:
    def __init__(self) -> None:
        self.started: list[tuple[str, str]] = []

    def start_task(self, task_id: str, task_type: str) -> None:
        self.started.append((task_id, task_type))


class FailingTaskWorkflowClient:
    def start_task(self, task_id: str, task_type: str) -> None:
        raise RuntimeError("temporal unavailable")


def make_client(
    tmp_path: Path,
    workflow_client: FakeTaskWorkflowClient | FailingTaskWorkflowClient | None = None,
    *,
    raise_server_exceptions: bool = True,
    access_control_enabled: bool = True,
) -> tuple[TestClient, FakeTaskWorkflowClient | FailingTaskWorkflowClient]:
    reset_engine_cache()
    settings = Settings(
        app_env="test",
        database_url=f"sqlite+pysqlite:///{tmp_path / 'tasks-api.db'}",
        pass_phrase_hash=hash_passphrase("weekly-secret"),
        pass_phrase_version="2026-W17",
        access_control_enabled=access_control_enabled,
        task_center_create_enabled=True,
        session_secret="test-secret",
        artifact_root=str(tmp_path / "artifacts"),
        corpus_root=str(tmp_path / "corpus"),
        workspace_root="/Users/xyc/Documents/codexwork",
    )
    app = create_app(settings)
    fake_workflow = workflow_client or FakeTaskWorkflowClient()
    app.dependency_overrides[get_task_workflow_client] = lambda: fake_workflow
    return TestClient(app, raise_server_exceptions=raise_server_exceptions), fake_workflow


def authenticate(client: TestClient) -> None:
    response = client.post("/api/access/verify", json={"passphrase": "weekly-secret"})
    assert response.status_code == 200


def token_from_url(url: str, name: str) -> str:
    values = parse_qs(urlparse(url).query).get(name)
    assert values
    return values[0]


def create_single_task(client: TestClient) -> dict:
    authenticate(client)
    response = client.post(
        "/api/tasks",
        json={"task_type": "single", "vehicles": [{"query": "风云X3 PLUS"}]},
    )
    assert response.status_code == 200
    return response.json()


def test_post_tasks_creates_single_vehicle_task_and_returns_access_urls(tmp_path: Path) -> None:
    client, fake_workflow = make_client(tmp_path)

    unauthorized = client.post(
        "/api/tasks",
        json={"task_type": "single", "vehicles": [{"query": "风云X3 PLUS"}]},
    )
    assert unauthorized.status_code == 401

    payload = create_single_task(client)

    assert payload["task_id"].startswith("task_")
    assert payload["status"] == "queued"
    assert payload["view_url"].startswith(f"/tasks/{payload['task_id']}?view_token=")
    assert payload["manage_url"].startswith(f"/tasks/{payload['task_id']}/manage?manage_token=")
    assert fake_workflow.started == [(payload["task_id"], "single")]

    view_token = token_from_url(payload["view_url"], "view_token")
    manage_token = token_from_url(payload["manage_url"], "manage_token")
    assert view_token != manage_token

    session = get_session_local()()
    try:
        task = session.get(Task, payload["task_id"])
        assert task is not None
        assert task.display_name == "风云X3 PLUS"
        assert task.task_type == "single"
        assert task.status == "queued"
        assert task.current_stage == "queued"
        assert task.view_token_hash == hash_task_token(view_token)
        assert task.manage_token_hash == hash_task_token(manage_token)
        assert view_token not in task.view_token_hash
        assert manage_token not in task.manage_token_hash
        assert len(task.vehicles) == 1
        assert task.vehicles[0].query == "风云X3 PLUS"
        assert task.vehicles[0].model_name == "风云X3 PLUS"
        assert [event.event_type for event in task.events] == ["created"]
    finally:
        session.close()


def test_post_tasks_returns_503_when_task_center_creation_disabled(tmp_path: Path) -> None:
    reset_engine_cache()
    settings = Settings(
        app_env="test",
        database_url=f"sqlite+pysqlite:///{tmp_path / 'tasks-disabled.db'}",
        pass_phrase_hash=hash_passphrase("weekly-secret"),
        pass_phrase_version="2026-W17",
        access_control_enabled=True,
        task_center_create_enabled=False,
        session_secret="test-secret",
        artifact_root=str(tmp_path / "artifacts"),
        workspace_root="/Users/xyc/Documents/codexwork",
    )
    app = create_app(settings)
    fake_workflow = FakeTaskWorkflowClient()
    app.dependency_overrides[get_task_workflow_client] = lambda: fake_workflow
    client = TestClient(app)
    authenticate(client)

    response = client.post(
        "/api/tasks",
        json={"task_type": "single", "vehicles": [{"query": "风云X3 PLUS"}]},
    )

    assert response.status_code == 503
    assert response.json()["detail"] == "task center creation is temporarily disabled; use the vehicle query flow"
    assert fake_workflow.started == []

    session = get_session_local()()
    try:
        assert session.query(Task).count() == 0
        assert session.query(TaskEvent).count() == 0
    finally:
        session.close()


def test_post_tasks_allows_direct_access_when_access_control_disabled(tmp_path: Path) -> None:
    client, fake_workflow = make_client(tmp_path, access_control_enabled=False)

    response = client.post(
        "/api/tasks",
        json={"task_type": "single", "vehicles": [{"query": "风云X3 PLUS"}]},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "queued"
    assert fake_workflow.started == [(payload["task_id"], "single")]

    list_response = client.get("/api/tasks")
    assert list_response.status_code == 200
    assert list_response.json()[0]["task_id"] == payload["task_id"]


def test_get_tasks_returns_list_after_passphrase_access(tmp_path: Path) -> None:
    client, _ = make_client(tmp_path)
    create_payload = create_single_task(client)

    unauthenticated_client, _ = make_client(tmp_path)
    unauthorized = unauthenticated_client.get("/api/tasks")
    assert unauthorized.status_code == 401

    response = client.get("/api/tasks")

    assert response.status_code == 200
    items = response.json()
    assert len(items) == 1
    assert items[0]["task_id"] == create_payload["task_id"]
    assert items[0]["display_name"] == "风云X3 PLUS"
    assert "view_token_hash" not in items[0]
    assert "manage_token_hash" not in items[0]


def test_get_tasks_load_returns_task_load_projection(tmp_path: Path) -> None:
    client, _ = make_client(tmp_path)
    create_single_task(client)

    response = client.get("/api/tasks/load")

    assert response.status_code == 200
    payload = response.json()
    assert payload["running_task_count"] == 0
    assert payload["queued_task_count"] == 1
    assert isinstance(payload["platforms"], dict)


def test_post_tasks_marks_task_failed_when_workflow_start_fails(tmp_path: Path) -> None:
    client, _ = make_client(
        tmp_path,
        workflow_client=FailingTaskWorkflowClient(),
        raise_server_exceptions=False,
    )
    authenticate(client)

    response = client.post(
        "/api/tasks",
        json={"task_type": "single", "vehicles": [{"query": "风云X3 PLUS"}]},
    )

    assert response.status_code == 503
    assert response.json()["detail"] == "task workflow unavailable"

    session = get_session_local()()
    try:
        tasks = session.query(Task).all()
        assert len(tasks) == 1
        task = tasks[0]
        assert task.status == "failed"
        assert task.current_stage == "workflow_start_failed"
        events = [event.event_type for event in task.events]
        assert events == ["created", "workflow_start_failed"]
        failure_event = task.events[-1]
        assert failure_event.payload_json == {"error": "temporal unavailable"}
    finally:
        session.close()


def test_get_task_detail_accepts_view_token_without_passphrase(tmp_path: Path) -> None:
    client, _ = make_client(tmp_path)
    payload = create_single_task(client)
    view_token = token_from_url(payload["view_url"], "view_token")

    public_client, _ = make_client(tmp_path)
    unauthorized = public_client.get(f"/api/tasks/{payload['task_id']}")
    assert unauthorized.status_code == 401

    wrong_token = public_client.get(f"/api/tasks/{payload['task_id']}?view_token=wrong")
    assert wrong_token.status_code == 403

    response = public_client.get(f"/api/tasks/{payload['task_id']}?view_token={view_token}")

    assert response.status_code == 200
    detail = response.json()
    assert detail["task_id"] == payload["task_id"]
    assert detail["vehicles"][0]["query"] == "风云X3 PLUS"
    assert detail["events"][0]["event_type"] == "created"
    assert "view_token_hash" not in detail
    assert "manage_token_hash" not in detail


def test_task_artifact_download_serves_job_and_corpus_files_with_view_token(tmp_path: Path) -> None:
    client, _ = make_client(tmp_path)
    payload = create_single_task(client)
    view_token = token_from_url(payload["view_url"], "view_token")

    job_file = tmp_path / "artifacts" / payload["task_id"] / "outputs" / "ai" / "final_report.json"
    corpus_file = tmp_path / "corpus" / "风云X3 PLUS__autohome-8089__dcd-25398" / "autohome" / "raw.xlsx"
    job_file.parent.mkdir(parents=True, exist_ok=True)
    corpus_file.parent.mkdir(parents=True, exist_ok=True)
    job_file.write_text('{"ok": true}', encoding="utf-8")
    corpus_file.write_bytes(b"excel-bytes")

    session = get_session_local()()
    try:
        job_artifact = TaskArtifact(task_id=payload["task_id"], artifact_type="json", path=str(job_file), downloadable=True)
        corpus_artifact = TaskArtifact(task_id=payload["task_id"], artifact_type="excel", path=str(corpus_file), downloadable=True)
        session.add_all([job_artifact, corpus_artifact])
        session.commit()
        job_artifact_id = job_artifact.id
        corpus_artifact_id = corpus_artifact.id
    finally:
        session.close()

    public_client, _ = make_client(tmp_path)
    missing_token = public_client.get(f"/api/tasks/{payload['task_id']}/artifacts/{job_artifact_id}")
    assert missing_token.status_code == 401

    wrong_token = public_client.get(f"/api/tasks/{payload['task_id']}/artifacts/{job_artifact_id}?view_token=wrong")
    assert wrong_token.status_code == 403

    job_response = public_client.get(f"/api/tasks/{payload['task_id']}/artifacts/{job_artifact_id}?view_token={view_token}")
    assert job_response.status_code == 200
    assert job_response.content == b'{"ok": true}'
    assert "attachment" in job_response.headers.get("content-disposition", "").lower()

    corpus_response = public_client.get(f"/api/tasks/{payload['task_id']}/artifacts/{corpus_artifact_id}?view_token={view_token}")
    assert corpus_response.status_code == 200
    assert corpus_response.content == b"excel-bytes"


def _write_summary_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "总览摘要"
    overview.append(["模块", "内容"])
    overview.append(["平台样本", "汽车之家 100 条 / 懂车帝 143 条"])
    compare = workbook.create_sheet("跨平台对比")
    compare.append(["维度", "汽车之家", "懂车帝"])
    compare.append(["空间", "正向集中", "正向集中"])
    business = workbook.create_sheet("综合业务摘要")
    business.append(["模块", "结论"])
    business.append(["核心卖点", "空间和配置好评突出"])
    opportunities = workbook.create_sheet("产品机会点")
    opportunities.append(["机会", "建议"])
    opportunities.append(["车机", "优先治理偶发卡顿"])
    one_pager = workbook.create_sheet("一页纸总结")
    one_pager.append(["双平台口碑一页纸总结"])
    one_pager.append(["空间好评突出，车机体验需要优化。"])
    workbook.save(path)


def _write_terms_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    positive = workbook.active
    positive.title = "positive_terms"
    positive.append(["term", "weight"])
    positive.append(["空间", 12])
    negative = workbook.create_sheet("negative_terms")
    negative.append(["term", "weight"])
    negative.append(["车机", 8])
    workbook.save(path)


def test_task_result_endpoint_assembles_one_pager_and_zip_with_view_token(tmp_path: Path) -> None:
    client, _ = make_client(tmp_path)
    payload = create_single_task(client)
    view_token = token_from_url(payload["view_url"], "view_token")

    output_root = tmp_path / "artifacts" / payload["task_id"] / "outputs"
    summary_path = output_root / "summary" / "风云X3 PLUS_双平台口碑摘要.xlsx"
    terms_path = output_root / "wordcloud" / "风云X3 PLUS_词云词项清单.xlsx"
    positive_png = output_root / "wordcloud" / "风云X3 PLUS_优点词云.png"
    final_report = output_root / "ai" / "final_report.json"
    analysis_facts = output_root / "ai" / "analysis_facts.jsonl"
    pdf_path = output_root / "report" / "风云X3 PLUS_完整报告.pdf"
    raw_path = tmp_path / "corpus" / "风云X3 PLUS__autohome-8089__dcd-25398" / "autohome" / "raw.xlsx"
    _write_summary_workbook(summary_path)
    _write_terms_workbook(terms_path)
    positive_png.parent.mkdir(parents=True, exist_ok=True)
    positive_png.write_bytes(b"png")
    final_report.parent.mkdir(parents=True, exist_ok=True)
    final_report.write_text(
        '{"headline":"风云X3 PLUS 空间好评突出","executive_summary":"基于双平台样本生成。",'
        '"strength_blocks":[{"title":"核心好评","summary":"空间表现突出","evidence_ids":["autohome_0001"]}],'
        '"weakness_blocks":[{"title":"核心槽点","summary":"车机偶发卡顿","evidence_ids":["dcd_0001"]}],'
        '"action_blocks":[{"title":"产品建议","summary":"优先治理车机稳定性","evidence_ids":["dcd_0001"]}],'
        '"boss_brief":["空间可作为传播主线","车机体验需要跟进"]}',
        encoding="utf-8",
    )
    analysis_facts.write_text(
        '{"comment_id":"autohome_0001","platform":"汽车之家","full_text":"空间很大，第三排够用。"}\n'
        '{"comment_id":"dcd_0001","platform":"懂车帝","full_text":"车机偶发卡顿，但动力顺。"}\n',
        encoding="utf-8",
    )
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    pdf_path.write_bytes(b"%PDF-1.4 report")
    raw_path.parent.mkdir(parents=True, exist_ok=True)
    raw_path.write_bytes(b"raw-excel")

    session = get_session_local()()
    try:
        task = session.get(Task, payload["task_id"])
        assert task is not None
        task.status = "completed"
        task.current_stage = "completed"
        task.completed_at = task.created_at
        task.vehicles[0].model_name = "风云X3 PLUS"
        task.vehicles[0].autohome_series_id = "8089"
        task.vehicles[0].dcd_series_id = "25398"
        session.add_all(
            [
                TaskArtifact(task_id=payload["task_id"], artifact_type="excel", path=str(summary_path), downloadable=True),
                TaskArtifact(task_id=payload["task_id"], artifact_type="excel", path=str(terms_path), downloadable=True),
                TaskArtifact(task_id=payload["task_id"], artifact_type="image_png", path=str(positive_png), downloadable=True),
                TaskArtifact(task_id=payload["task_id"], artifact_type="json", path=str(final_report), downloadable=True),
                TaskArtifact(task_id=payload["task_id"], artifact_type="jsonl", path=str(analysis_facts), downloadable=True),
                TaskArtifact(task_id=payload["task_id"], artifact_type="pdf", path=str(pdf_path), downloadable=True),
                TaskArtifact(task_id=payload["task_id"], artifact_type="excel", path=str(raw_path), downloadable=True),
            ]
        )
        session.commit()
    finally:
        session.close()

    public_client, _ = make_client(tmp_path)
    unauthorized = public_client.get(f"/api/tasks/{payload['task_id']}/result")
    assert unauthorized.status_code == 401

    response = public_client.get(f"/api/tasks/{payload['task_id']}/result?view_token={view_token}")
    assert response.status_code == 200
    result = response.json()
    assert result["task_id"] == payload["task_id"]
    assert result["report_ready"] is True
    assert result["sample_summary"] == {"autohome_count": 100, "dcd_count": 143}
    assert result["ai_report"]["headline"] == "风云X3 PLUS 空间好评突出"
    assert result["template_report"]["title"] == "双平台口碑一页纸总结"
    assert result["wordcloud"]["keyword_rankings"]["positive"][0] == {"term": "空间", "count": 12}
    assert result["evidence_samples"][0]["comment_id"] == "autohome_0001"
    assert result["zip_url"].endswith(f"/api/tasks/{payload['task_id']}/artifacts.zip")

    zip_response = public_client.get(f"/api/tasks/{payload['task_id']}/artifacts.zip?view_token={view_token}")
    assert zip_response.status_code == 200
    archive_path = tmp_path / "result.zip"
    archive_path.write_bytes(zip_response.content)
    with zipfile.ZipFile(archive_path) as archive:
        names = sorted(archive.namelist())
    assert "风云X3 PLUS_完整报告.pdf" in names
    assert "raw/autohome_raw.xlsx" in names
    assert "ai/final_report.json" in names


def test_management_action_requires_passphrase_session_and_manage_token(tmp_path: Path) -> None:
    client, _ = make_client(tmp_path)
    payload = create_single_task(client)
    manage_token = token_from_url(payload["manage_url"], "manage_token")

    public_client, _ = make_client(tmp_path)
    missing_session = public_client.post(f"/api/tasks/{payload['task_id']}/cancel?manage_token={manage_token}")
    assert missing_session.status_code == 401

    missing_token = client.post(f"/api/tasks/{payload['task_id']}/cancel")
    assert missing_token.status_code == 403

    wrong_token = client.post(f"/api/tasks/{payload['task_id']}/cancel?manage_token=wrong")
    assert wrong_token.status_code == 403

    response = client.post(f"/api/tasks/{payload['task_id']}/cancel?manage_token={manage_token}")

    assert response.status_code == 200
    detail = response.json()
    assert detail["status"] == "cancelled"
    assert detail["current_stage"] == "cancelled"
    assert detail["events"][-1]["event_type"] == "cancel_requested"

    retry_response = client.post(f"/api/tasks/{payload['task_id']}/retry?manage_token={manage_token}")
    assert retry_response.status_code == 200
    assert retry_response.json()["status"] == "queued"
    assert retry_response.json()["events"][-1]["event_type"] == "manual_retry_requested"

    pause_response = client.post(f"/api/tasks/{payload['task_id']}/pause-retry?manage_token={manage_token}")
    assert pause_response.status_code == 200
    pause_detail = pause_response.json()
    assert pause_detail["status"] == "retry_paused"
    assert pause_detail["current_stage"] == "retry_paused"
    assert pause_detail["events"][-1]["event_type"] == "retry_paused"

    session = get_session_local()()
    try:
        events = (
            session.query(TaskEvent)
            .filter(TaskEvent.task_id == payload["task_id"])
            .order_by(TaskEvent.id)
            .all()
        )
        assert [event.event_type for event in events] == [
            "created",
            "cancel_requested",
            "manual_retry_requested",
            "retry_paused",
        ]
    finally:
        session.close()
