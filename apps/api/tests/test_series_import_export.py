from __future__ import annotations

from io import BytesIO
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.config import Settings
from app.db import init_db, reset_engine_cache
from app.models import ConfirmedVehicleSeries, SeriesConflict, SeriesImportBatch
from app.services.confirmed_vehicle_series import query_key
from app.services.series_admin import SeriesMutation, create_series_record
from app.services.series_import_export import commit_series_import, export_series_excel, preview_series_import


def _session(tmp_path: Path):
    reset_engine_cache()
    settings = Settings(app_env="test", database_url=f"sqlite+pysqlite:///{tmp_path / 'series.db'}")
    init_db(settings)
    engine = create_engine(settings.database_url, future=True)
    return sessionmaker(bind=engine, future=True)


def _csv(rows: list[dict[str, str]]) -> bytes:
    header = ["query", "platform", "series_id", "url", "title", "source"]
    lines = [",".join(header)]
    for row in rows:
        lines.append(",".join(row.get(column, "") for column in header))
    return ("\ufeff" + "\n".join(lines) + "\n").encode("utf-8")


def _xlsx(rows: list[dict[str, str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["query", "platform", "series_id", "url", "title", "source"])
    for row in rows:
        ws.append([row.get("query"), row.get("platform"), row.get("series_id"), row.get("url"), row.get("title"), row.get("source")])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _seed_existing(db) -> None:
    create_series_record(
        db,
        SeriesMutation(
            query="风云T11",
            platform="autohome",
            series_id="7411",
            url="https://k.autohome.com.cn/7411",
            title="风云T11",
            source="manual",
            operator="tester",
            reason="seed",
        ),
    )
    db.commit()


def test_preview_counts_new_duplicate_and_conflict(tmp_path: Path) -> None:
    SessionLocal = _session(tmp_path)
    with SessionLocal() as db:
        _seed_existing(db)

        preview = preview_series_import(
            db,
            filename="series.csv",
            content=_csv(
                [
                    {"query": "风云T11", "platform": "autohome", "series_id": "7411"},
                    {"query": "风云T11", "platform": "autohome", "series_id": "9999"},
                    {"query": "风云T11", "platform": "dongchedi", "series_id": "5498"},
                ]
            ),
        )

        assert preview.summary == {"new": 1, "duplicate": 1, "conflict": 1, "invalid": 0}
        assert [row.status for row in preview.rows] == ["duplicate", "conflict", "new"]
        conflict = preview.rows[1]
        assert conflict.existing_value_json["series_id"] == "7411"
        assert conflict.incoming_value_json["series_id"] == "9999"


def test_commit_import_writes_only_new_rows_and_conflicts(tmp_path: Path) -> None:
    SessionLocal = _session(tmp_path)
    with SessionLocal() as db:
        _seed_existing(db)

        preview = commit_series_import(
            db,
            filename="series.csv",
            operator="operator-1",
            content=_csv(
                [
                    {"query": "风云T11", "platform": "autohome", "series_id": "7411"},
                    {"query": "风云T11", "platform": "autohome", "series_id": "9999", "url": "https://example.test/9999", "title": "T11 alt", "source": "sheet"},
                    {"query": "风云T11", "platform": "dongchedi", "series_id": "5498", "url": "https://example.test/5498", "title": "T11 DCD", "source": "sheet"},
                ]
            ),
        )

        assert preview.summary == {"new": 1, "duplicate": 1, "conflict": 1, "invalid": 0}
        batch = db.query(SeriesImportBatch).one()
        assert batch.source == "import"
        assert batch.filename == "series.csv"
        assert batch.operator == "operator-1"
        assert batch.summary_json == preview.summary
        records = db.query(ConfirmedVehicleSeries).order_by(ConfirmedVehicleSeries.platform).all()
        assert [(row.platform, row.series_id, row.import_batch_id) for row in records] == [
            ("autohome", "7411", None),
            ("dongchedi", "5498", batch.id),
        ]
        conflicts = db.query(SeriesConflict).all()
        assert len(conflicts) == 1
        assert conflicts[0].conflict_type == "series_import"
        assert conflicts[0].status == "open"
        assert conflicts[0].import_batch_id == batch.id
        assert conflicts[0].existing_value_json["series_id"] == "7411"
        assert conflicts[0].incoming_value_json["series_id"] == "9999"
        assert conflicts[0].incoming_value_json["url"] == "https://example.test/9999"


def test_export_series_excel_contains_active_rows(tmp_path: Path) -> None:
    SessionLocal = _session(tmp_path)
    with SessionLocal() as db:
        _seed_existing(db)

        content = export_series_excel(db)

    wb = load_workbook(BytesIO(content))
    ws = wb["confirmed_vehicle_series"]
    assert [cell.value for cell in ws[1]] == ["query", "platform", "series_id", "url", "title", "source", "status", "updated_at"]
    assert [cell.value for cell in ws[2][:7]] == ["风云T11", "autohome", "7411", "https://k.autohome.com.cn/7411", "风云T11", "manual", "active"]


def test_preview_supports_xlsx_with_same_columns(tmp_path: Path) -> None:
    SessionLocal = _session(tmp_path)
    with SessionLocal() as db:
        preview = preview_series_import(
            db,
            filename="series.xlsx",
            content=_xlsx([{"query": "风云T11", "platform": "autohome", "series_id": "7411"}]),
        )

        assert preview.summary == {"new": 1, "duplicate": 0, "conflict": 0, "invalid": 0}
        assert preview.rows[0].query_key == query_key("风云T11")


def test_invalid_rows_are_counted_but_not_committed(tmp_path: Path) -> None:
    SessionLocal = _session(tmp_path)
    with SessionLocal() as db:
        preview = commit_series_import(
            db,
            filename="series.csv",
            operator="operator-1",
            content=_csv(
                [
                    {"query": "", "platform": "autohome", "series_id": "7411"},
                    {"query": "风云T11", "platform": "bad-platform", "series_id": "7411"},
                    {"query": "风云T11", "platform": "autohome", "series_id": ""},
                ]
            ),
        )

        assert preview.summary == {"new": 0, "duplicate": 0, "conflict": 0, "invalid": 3}
        assert db.query(ConfirmedVehicleSeries).count() == 0
        assert db.query(SeriesConflict).count() == 0
        assert db.query(SeriesImportBatch).one().summary_json == preview.summary


def test_duplicate_rows_are_counted_but_not_stored_as_conflicts(tmp_path: Path) -> None:
    SessionLocal = _session(tmp_path)
    with SessionLocal() as db:
        _seed_existing(db)

        preview = commit_series_import(
            db,
            filename="series.csv",
            operator="operator-1",
            content=_csv([{"query": "风云T11", "platform": "autohome", "series_id": "7411"}]),
        )

        assert preview.summary == {"new": 0, "duplicate": 1, "conflict": 0, "invalid": 0}
        assert db.query(ConfirmedVehicleSeries).count() == 1
        assert db.query(SeriesConflict).count() == 0
