from __future__ import annotations

from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.config import Settings
from app.db import init_db, reset_engine_cache
from app.models import Base, ComparisonJob, ConfirmedVehicleSeries, Job, KoubeiRawComment


def _seed_db(database_path: Path, *, include_test_only: bool) -> str:
    database_url = f"sqlite+pysqlite:///{database_path}"
    reset_engine_cache()
    init_db(Settings(app_env="test", database_url=database_url))
    engine = create_engine(database_url, future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, future=True)
    with SessionLocal() as db:
        db.add(
            ConfirmedVehicleSeries(
                query_key="风云t11",
                query="风云T11",
                platform="autohome",
                series_id="7411",
                status="active",
            )
        )
        if include_test_only:
            db.add(
                ConfirmedVehicleSeries(
                    query_key="风云t11",
                    query="风云T11",
                    platform="dongchedi",
                    series_id="9436",
                    status="active",
                )
            )
            db.add(
                KoubeiRawComment(
                    query_key="风云t11",
                    query="风云T11",
                    model_name="风云T11",
                    platform="autohome",
                    series_id="7411",
                    source_link="https://example.test/comment/1",
                    dedupe_key="link:https://example.test/comment/1",
                    row_json={"content": "新增评论"},
                )
            )
            db.add(Job(job_id="job_test_001", query="风云T11", model_name="风云T11", status="completed", passphrase_version="dev"))
            db.add(ComparisonJob(comparison_id="cmp_test_001", status="completed", vehicle_count=2, passphrase_version="dev"))
        db.commit()
    return database_url


def test_cutover_audit_exports_series_raw_comment_and_task_sheets(tmp_path: Path) -> None:
    prod_url = _seed_db(tmp_path / "prod.db", include_test_only=False)
    test_url = _seed_db(tmp_path / "test.db", include_test_only=True)
    output = tmp_path / "cutover-audit.xlsx"

    result = subprocess.run(
        [
            sys.executable,
            "scripts/series-admin/audit_cutover.py",
            "--prod-url",
            prod_url,
            "--test-url",
            test_url,
            "--output",
            str(output),
        ],
        cwd=Path(__file__).resolve().parents[3],
        text=True,
        capture_output=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert "series_new=1" in result.stdout
    assert "raw_comments_new=1" in result.stdout
    assert "task_results_new=2" in result.stdout
    workbook = load_workbook(output)
    assert {"summary", "series_diff", "raw_comment_diff", "task_result_diff"}.issubset(set(workbook.sheetnames))
