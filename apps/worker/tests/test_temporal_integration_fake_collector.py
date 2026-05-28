from __future__ import annotations

import asyncio
from collections import Counter
from datetime import UTC, datetime
import json
import os
import sys
from pathlib import Path
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, text

WORKER_ROOT = Path(__file__).resolve().parents[1]
APPS_ROOT = Path(__file__).resolve().parents[2]
API_ROOT = APPS_ROOT / "api"
for root in (WORKER_ROOT, API_ROOT):
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))

os.environ["DATABASE_URL"] = "sqlite+pysqlite:///:memory:"

from app.config import Settings
from app.db import get_session_local, reset_engine_cache
from app.main import create_app
from app.models import Task
from app.services.passphrase import hash_passphrase
from app.services.task_workflow_client import get_task_workflow_client
from worker_app.corpus import AUTOHOME_HEADERS, DCD_HEADERS, load_platform_state
from worker_app.task_eta import estimate_queue_seconds, eta_reason_for_platform
from worker_app.task_scheduler import QueuedRun, plan_dispatch_order
from worker_app.task_store import TaskStore, utc_now_iso
from worker_app.temporal_activities import TaskActivities
from worker_app.temporal_workflows import run_comparison_task, run_single_vehicle_task


class FakeTaskWorkflowClient:
    def __init__(self) -> None:
        self.started: list[tuple[str, str]] = []

    def start_task(self, task_id: str, task_type: str) -> None:
        self.started.append((task_id, task_type))


def _make_client(tmp_path: Path) -> tuple[TestClient, FakeTaskWorkflowClient, str]:
    reset_engine_cache()
    database_url = f"sqlite+pysqlite:///{tmp_path / 'integration.db'}"
    settings = Settings(
        app_env="test",
        database_url=database_url,
        pass_phrase_hash=hash_passphrase("weekly-secret"),
        pass_phrase_version="2026-W17",
        task_center_create_enabled=True,
        session_secret="test-secret",
        artifact_root=str(tmp_path / "artifacts"),
        workspace_root=str(tmp_path),
    )
    app = create_app(settings)
    workflow_client = FakeTaskWorkflowClient()
    app.dependency_overrides[get_task_workflow_client] = lambda: workflow_client
    return TestClient(app), workflow_client, database_url


def _authenticate(client: TestClient) -> None:
    response = client.post("/api/access/verify", json={"passphrase": "weekly-secret"})
    assert response.status_code == 200


def _create_task(client: TestClient, task_type: str, vehicles: list[str]) -> dict[str, Any]:
    response = client.post(
        "/api/tasks",
        json={"task_type": task_type, "vehicles": [{"query": vehicle} for vehicle in vehicles]},
    )
    assert response.status_code == 200, response.text
    return response.json()


def _set_vehicle_series(database_url: str, task_id: str, series_by_query: dict[str, tuple[str, str]]) -> None:
    engine = create_engine(database_url, future=True, connect_args={"check_same_thread": False})
    with engine.begin() as connection:
        rows = connection.execute(
            text("SELECT id, query FROM task_vehicles WHERE task_id = :task_id ORDER BY position ASC"),
            {"task_id": task_id},
        ).mappings().all()
        for row in rows:
            autohome_series_id, dcd_series_id = series_by_query[str(row["query"])]
            connection.execute(
                text(
                    """
                    UPDATE task_vehicles
                    SET model_name = query,
                        autohome_series_id = :autohome_series_id,
                        dcd_series_id = :dcd_series_id,
                        updated_at = :updated_at
                    WHERE id = :vehicle_id
                    """
                ),
                {
                    "vehicle_id": int(row["id"]),
                    "autohome_series_id": autohome_series_id,
                    "dcd_series_id": dcd_series_id,
                    "updated_at": utc_now_iso(),
                },
            )
    engine.dispose()


def _seed_confirmed_series(database_url: str, query: str, autohome_series_id: str, dcd_series_id: str) -> None:
    engine = create_engine(database_url, future=True, connect_args={"check_same_thread": False})
    now = utc_now_iso()
    query_key = query.strip().lower()
    with engine.begin() as connection:
        for platform, series_id in (("autohome", autohome_series_id), ("dongchedi", dcd_series_id)):
            connection.execute(
                text(
                    """
                    INSERT INTO confirmed_vehicle_series (
                        query_key, query, platform, series_id, status, created_at, updated_at
                    )
                    VALUES (
                        :query_key, :query, :platform, :series_id, 'active', :created_at, :updated_at
                    )
                    """
                ),
                {
                    "query_key": query_key,
                    "query": query,
                    "platform": platform,
                    "series_id": series_id,
                    "created_at": now,
                    "updated_at": now,
                },
            )
    engine.dispose()


def _write_platform_workbook(path: Path, *, platform: str, model_name: str, run_id: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "口碑明细"
    headers = AUTOHOME_HEADERS if platform == "autohome" else DCD_HEADERS
    sheet.append(headers)
    if platform == "autohome":
        row = {
            "用户名": f"{model_name}车主A",
            "发表日期": "2026-05-01",
            "车型": model_name,
            "评价详情": f"{model_name} 汽车之家 fake collector 评论",
            "来源链接": f"https://k.autohome.example/{run_id}",
            "抓取页码": "1",
        }
    else:
        row = {
            "用户名": f"{model_name}车主B",
            "发布时间": "2026-05-02",
            "评价车型": model_name,
            "评价全文": f"{model_name} 懂车帝 fake collector 评论",
            "来源链接": f"https://dongchedi.example/{run_id}",
            "抓取页码": "1",
        }
    sheet.append([row.get(header, "") for header in headers])
    workbook.save(path)


def _task_model_name(payload: dict[str, Any]) -> str:
    task = payload.get("task") if isinstance(payload, dict) else {}
    vehicles = task.get("vehicles") if isinstance(task, dict) else []
    if vehicles and isinstance(vehicles[0], dict):
        return str(vehicles[0].get("model_name") or vehicles[0].get("query") or "车型A")
    return str(task.get("display_name") or "车型A") if isinstance(task, dict) else "车型A"


def _write_fake_summary_workbook(path: Path, model_name: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "总览摘要"
    overview.append(["模块", "内容"])
    overview.append(["平台样本", "汽车之家 1 条 / 懂车帝 1 条"])
    compare = workbook.create_sheet("跨平台对比")
    compare.append(["维度", "汽车之家", "懂车帝"])
    compare.append(["空间", "正向集中", "正向集中"])
    business = workbook.create_sheet("综合业务摘要")
    business.append(["模块", "结论"])
    business.append(["核心卖点", f"{model_name} 双平台评论样本稳定"])
    opportunities = workbook.create_sheet("产品机会点")
    opportunities.append(["机会", "建议"])
    opportunities.append(["车机", "持续观察新增负向评论"])
    one_pager = workbook.create_sheet("一页纸总结")
    one_pager.append([f"{model_name} 双平台口碑一页纸总结"])
    one_pager.append(["当前 fake collector 样本已生成完整智能报告产物。"])
    workbook.save(path)


def _write_fake_terms_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    positive = workbook.active
    positive.title = "positive_terms"
    positive.append(["term", "weight"])
    positive.append(["空间", 12])
    negative = workbook.create_sheet("negative_terms")
    negative.append(["term", "weight"])
    negative.append(["车机", 8])
    workbook.save(path)


class FakeCollectorActivityRunner:
    def __init__(self, database_url: str, output_root: Path, *, fail_dongchedi_once: bool = False) -> None:
        self.database_url = database_url
        self.activities = TaskActivities(database_url)
        self.output_root = output_root
        self.fail_dongchedi_once = fail_dongchedi_once
        self.failed_dongchedi = False
        self.calls: list[str] = []

    async def __call__(self, name: str, payload: Any, timeout: Any) -> dict[str, Any]:
        self.calls.append(name)
        if name == "create_or_join_collection_run":
            run = await self.activities.create_or_join_collection_run(payload)
            return run
        if name == "dispatch_collection_run":
            run = TaskStore(self.database_url).load_collection_run(str(payload["run_id"]))
            run_payload = _run_to_dict(run)
            if run_payload["platform"] == "dongchedi" and self.fail_dongchedi_once and not self.failed_dongchedi:
                self.failed_dongchedi = True
                return self._finish_run(run_payload, status="failed", failure_category="timeout")
            else:
                return self._finish_run(run_payload, status="completed")
        if name == "retry_failed_platforms":
            return self._retry_failed_platforms(payload)
        if name == "run_postprocess":
            return self._write_fake_analysis_facts(payload)
        if name == "run_llm_report":
            return self._write_fake_report_outputs(payload)
        method = getattr(self.activities, name)
        return await method(payload)

    def _task_output_root(self, payload: dict[str, Any]) -> Path:
        task_id = str(payload.get("task_id") or "task")
        artifact_root = Path(os.environ.get("ARTIFACT_ROOT") or self.output_root)
        return artifact_root / task_id / "outputs"

    def _write_fake_analysis_facts(self, payload: dict[str, Any]) -> dict[str, Any]:
        model_name = _task_model_name(payload)
        output_root = self._task_output_root(payload) / "ai"
        output_root.mkdir(parents=True, exist_ok=True)
        facts_path = output_root / "analysis_facts.jsonl"
        facts_path.write_text(
            json.dumps(
                {
                    "comment_id": "fake-1",
                    "platform": "autohome",
                    "summary": f"{model_name} 空间表现获得正向反馈",
                },
                ensure_ascii=False,
            )
            + "\n",
            encoding="utf-8",
        )
        return {"artifact_paths": [str(facts_path)], "skipped": False}

    def _write_fake_report_outputs(self, payload: dict[str, Any]) -> dict[str, Any]:
        model_name = _task_model_name(payload)
        output_root = self._task_output_root(payload)
        ai_root = output_root / "ai"
        summary_root = output_root / "summary"
        wordcloud_root = output_root / "wordcloud"
        ai_root.mkdir(parents=True, exist_ok=True)
        wordcloud_root.mkdir(parents=True, exist_ok=True)
        report_path = ai_root / "final_report.json"
        metrics_path = ai_root / "llm_metrics.json"
        pdf_path = output_root / f"{model_name}_智能一页纸完整报告.pdf"
        summary_path = summary_root / f"{model_name}_双平台口碑摘要.xlsx"
        terms_path = wordcloud_root / f"{model_name}_词云词项清单.xlsx"
        positive_png = wordcloud_root / f"{model_name}_优点词云.png"
        negative_png = wordcloud_root / f"{model_name}_槽点词云.png"
        _write_fake_summary_workbook(summary_path, model_name)
        _write_fake_terms_workbook(terms_path)
        report_path.write_text(
            json.dumps(
                {
                    "headline": f"{model_name} 双平台口碑表现稳定",
                    "executive_summary": f"{model_name} 当前增量样本已完成智能报告生成。",
                    "key_findings": [
                        {
                            "title": "空间反馈稳定",
                            "summary": "正向评论集中在空间和配置。",
                            "evidence": ["车主反馈空间够用。"],
                        }
                    ],
                    "dimension_matrix": [
                        {"dimension": "空间", "sentiment": "positive", "summary": "正向集中"}
                    ],
                    "boss_brief": ["继续关注车机负向评论。"],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        metrics_path.write_text('{"mode":"fake"}', encoding="utf-8")
        pdf_path.write_bytes(b"%PDF-1.4\n% fake report\n")
        positive_png.write_bytes(b"fake-png")
        negative_png.write_bytes(b"fake-png")
        return {
            "artifact_paths": [
                str(report_path),
                str(metrics_path),
                str(pdf_path),
                str(summary_path),
                str(terms_path),
                str(positive_png),
                str(negative_png),
            ],
            "skipped": False,
        }

    def _finish_run(self, run: dict[str, Any], *, status: str, failure_category: str | None = None) -> dict[str, Any]:
        output_path: Path | None = None
        if status == "completed":
            output_path = self.output_root / str(run["platform"]) / f"{run['run_id']}.xlsx"
            _write_platform_workbook(
                output_path,
                platform=str(run["platform"]),
                model_name=str(run["model_name"]),
                run_id=str(run["run_id"]),
            )
        engine = create_engine(self.database_url, future=True, connect_args={"check_same_thread": False})
        with engine.begin() as connection:
            connection.execute(
                text(
                    """
                    UPDATE collection_runs
                    SET status = :status,
                        output_path = :output_path,
                        failure_category = :failure_category,
                        finished_at = :finished_at,
                        updated_at = :updated_at
                    WHERE run_id = :run_id
                    """
                ),
                {
                    "run_id": str(run["run_id"]),
                    "status": status,
                    "output_path": str(output_path) if output_path else None,
                    "failure_category": failure_category,
                    "finished_at": utc_now_iso(),
                    "updated_at": utc_now_iso(),
                },
            )
        engine.dispose()
        return _run_to_dict(TaskStore(self.database_url).load_collection_run(str(run["run_id"])))

    def _retry_failed_platforms(self, payload: dict[str, Any]) -> dict[str, Any]:
        TaskStore(self.database_url).append_task_event(str(payload["task_id"]), "retry_attempted", payload)
        successful_platforms: list[str] = []
        runs: dict[str, dict[str, Any]] = {}
        for failure in payload.get("failed_platforms", []):
            platform = str(failure.get("platform") or "")
            run_id = str(failure.get("run_id") or "")
            if platform != "dongchedi" or not run_id:
                continue
            run = TaskStore(self.database_url).load_collection_run(run_id)
            completed = self._finish_run(_run_to_dict(run), status="completed")
            successful_platforms.append(platform)
            runs[platform] = completed
        return {"successful_platforms": successful_platforms, "failed_platforms": [], "runs": runs}


def _run_to_dict(run: Any) -> dict[str, Any]:
    return {
        "run_id": run.run_id,
        "platform": run.platform,
        "query_key": run.query_key,
        "model_name": run.model_name,
        "series_id": run.series_id,
        "status": run.status,
        "mode": run.mode,
        "shared_by_task_ids": run.shared_by_task_ids,
        "failure_category": run.failure_category,
        "output_path": run.output_path,
    }


def test_api_created_single_task_runs_fake_collector_workflow_to_detail_payload(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("ARTIFACT_ROOT", str(tmp_path / "artifacts"))
    monkeypatch.setenv("KOUBEI_CORPUS_ROOT", str(tmp_path / "corpus"))
    client, workflow_client, database_url = _make_client(tmp_path)
    _authenticate(client)
    payload = _create_task(client, "single", ["车型A"])
    _set_vehicle_series(database_url, payload["task_id"], {"车型A": ("8089", "25398")})

    runner = FakeCollectorActivityRunner(database_url, tmp_path / "collector")
    result = asyncio.run(run_single_vehicle_task(payload["task_id"], runner))

    assert result == {"task_id": payload["task_id"], "status": "completed"}
    assert workflow_client.started == [(payload["task_id"], "single")]
    detail_response = client.get(f"/api/tasks/{payload['task_id']}")
    assert detail_response.status_code == 200
    detail = detail_response.json()
    assert detail["status"] in {"completed", "completed_degraded"}
    assert detail["artifacts"]
    assert any(event["event_type"] == "collection_run_shared" for event in detail["events"]) is False
    assert load_platform_state(database_url, query="车型A", platform="autohome", series_id="8089").existing_count == 1
    assert load_platform_state(database_url, query="车型A", platform="dongchedi", series_id="25398").existing_count == 1


def test_resolve_vehicle_inputs_backfills_missing_series_from_confirmed_table(tmp_path: Path) -> None:
    client, _workflow_client, database_url = _make_client(tmp_path)
    _authenticate(client)
    payload = _create_task(client, "single", ["车型A"])
    _seed_confirmed_series(database_url, "车型A", "8089", "25398")

    task = TaskStore(database_url).load_task(payload["task_id"])
    result = asyncio.run(
        TaskActivities(database_url=database_url).resolve_vehicle_inputs(
            {
                "task_id": payload["task_id"],
                "vehicles": [
                    {
                        "id": task.vehicles[0].id,
                        "query": task.vehicles[0].query,
                        "model_name": task.vehicles[0].model_name,
                        "autohome_series_id": task.vehicles[0].autohome_series_id,
                        "dcd_series_id": task.vehicles[0].dcd_series_id,
                    }
                ],
                "mode": "incremental",
            }
        )
    )

    assert result["ok"] is True
    assert result["platforms"]["autohome"]["series_id"] == "8089"
    assert result["platforms"]["dongchedi"]["series_id"] == "25398"

    updated = TaskStore(database_url).load_task(payload["task_id"])
    assert updated.vehicles[0].autohome_series_id == "8089"
    assert updated.vehicles[0].dcd_series_id == "25398"


def test_single_task_export_uses_backfilled_series_ids_for_corpus_dir(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("ARTIFACT_ROOT", str(tmp_path / "artifacts"))
    corpus_root = tmp_path / "corpus"
    monkeypatch.setenv("KOUBEI_CORPUS_ROOT", str(corpus_root))
    client, _workflow_client, database_url = _make_client(tmp_path)
    _authenticate(client)
    payload = _create_task(client, "single", ["车型A"])
    _seed_confirmed_series(database_url, "车型A", "8089", "25398")

    runner = FakeCollectorActivityRunner(database_url, tmp_path / "collector")
    result = asyncio.run(run_single_vehicle_task(payload["task_id"], runner))

    assert result == {"task_id": payload["task_id"], "status": "completed"}
    expected_dir = corpus_root / "车型A__autohome-8089__dcd-25398"
    assert (expected_dir / "manifest.json").exists()
    assert (expected_dir / "autohome" / "raw.xlsx").exists()
    assert (expected_dir / "dongchedi" / "raw.xlsx").exists()
    assert not (corpus_root / "车型A__autohome-vehicle__dcd-vehicle").exists()


def test_api_created_comparison_task_runs_child_task_workflows_to_task_result(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("ARTIFACT_ROOT", str(tmp_path / "artifacts"))
    monkeypatch.setenv("KOUBEI_CORPUS_ROOT", str(tmp_path / "corpus"))
    client, workflow_client, database_url = _make_client(tmp_path)
    _authenticate(client)
    payload = _create_task(client, "comparison", ["车型A", "车型B"])
    _seed_confirmed_series(database_url, "车型A", "8089", "25398")
    _seed_confirmed_series(database_url, "车型B", "8090", "25399")

    runner = FakeCollectorActivityRunner(database_url, tmp_path / "collector")
    child_tasks: dict[str, asyncio.Task] = {}

    async def child_workflow_runner(child_payload: dict[str, Any]) -> dict[str, Any]:
        child_task_id = str(child_payload["child_task_id"])
        child_tasks[child_task_id] = asyncio.create_task(run_single_vehicle_task(child_task_id, runner))
        return {"child_task_id": child_task_id, "started": True}

    async def child_completion_runner(child_payload: dict[str, Any]) -> dict[str, Any]:
        return await child_tasks[str(child_payload["child_task_id"])]

    result = asyncio.run(
        run_comparison_task(
            payload["task_id"],
            runner,
            child_workflow_runner=child_workflow_runner,
            child_completion_runner=child_completion_runner,
        )
    )

    assert result["status"] == "completed"
    assert workflow_client.started == [(payload["task_id"], "comparison")]
    detail_response = client.get(f"/api/tasks/{payload['task_id']}")
    assert detail_response.status_code == 200
    detail = detail_response.json()
    assert detail["status"] == "completed"
    assert len(detail["vehicles"]) == 2
    assert detail["artifacts"]


def test_four_user_fake_load_shares_runs_limits_lanes_and_records_upgrade(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("ARTIFACT_ROOT", str(tmp_path / "artifacts"))
    monkeypatch.setenv("KOUBEI_CORPUS_ROOT", str(tmp_path / "corpus"))
    client, _workflow_client, database_url = _make_client(tmp_path)
    _authenticate(client)
    series = {
        "车型A": ("8089", "25398"),
        "车型B": ("8090", "25399"),
        "车型C": ("8091", "25400"),
        "车型D": ("8092", "25401"),
        "车型E": ("8093", "25402"),
        "车型F": ("8094", "25403"),
        "车型G": ("8095", "25404"),
        "车型H": ("8096", "25405"),
    }
    task_1 = _create_task(client, "single", ["车型A"])["task_id"]
    task_2 = _create_task(client, "comparison", ["车型A", "车型B", "车型C"])["task_id"]
    task_3 = _create_task(client, "single", ["车型A"])["task_id"]
    task_4 = _create_task(client, "comparison", ["车型D", "车型E", "车型F", "车型G", "车型H"])["task_id"]
    for task_id in (task_1, task_2, task_3, task_4):
        _set_vehicle_series(database_url, task_id, series)

    store = TaskStore(database_url)
    session = get_session_local()()
    try:
        for task in session.query(Task).order_by(Task.created_at.asc(), Task.task_id.asc()).all():
            for vehicle in task.vehicles:
                store.create_or_join_collection_run(
                    platform="autohome",
                    query_key=vehicle.query,
                    model_name=vehicle.model_name,
                    series_id=str(vehicle.autohome_series_id),
                    mode="incremental",
                    task_id=task.task_id,
                )
                store.create_or_join_collection_run(
                    platform="dongchedi",
                    query_key=vehicle.query,
                    model_name=vehicle.model_name,
                    series_id=str(vehicle.dcd_series_id),
                    mode="incremental",
                    task_id=task.task_id,
                )
    finally:
        session.close()

    engine = create_engine(database_url, future=True, connect_args={"check_same_thread": False})
    with engine.begin() as connection:
        shared_a_runs = connection.execute(
            text(
                """
                SELECT platform, shared_by_task_ids
                FROM collection_runs
                WHERE query_key = '车型A'
                ORDER BY platform
                """
            )
        ).mappings().all()
        assert {row["platform"] for row in shared_a_runs} == {"autohome", "dongchedi"}
        for row in shared_a_runs:
            assert set(json.loads(row["shared_by_task_ids"])) == {task_1, task_2, task_3}

        task_types = dict(connection.execute(text("SELECT task_id, task_type FROM tasks")).all())
        platform_queues: dict[str, list[QueuedRun]] = {"autohome": [], "dongchedi": []}
        for row in connection.execute(
            text("SELECT run_id, platform, shared_by_task_ids FROM collection_runs WHERE status = 'queued'")
        ).mappings():
            owner_task_id = str(json.loads(row["shared_by_task_ids"])[0])
            platform_queues[str(row["platform"])].append(
                QueuedRun(
                    str(row["run_id"]),
                    owner_task_id,
                    str(task_types[owner_task_id]),
                    str(row["platform"]),
                    waited_seconds=400,
                )
            )

        decisions = plan_dispatch_order(
            platform_queues=platform_queues,
            agent_capacity={"autohome": 2, "dongchedi": 2},
            active_lanes_by_task={},
            single_task_priority_weight=0.1,
        )
        assert max(Counter(decision.task_id for decision in decisions).values()) == 1

        eta_reason = eta_reason_for_platform(
            "autohome",
            queue_length=len(platform_queues["autohome"]),
            agent_count=2,
            retrying_count=0,
        )
        for position, task_id in enumerate((task_1, task_2, task_3, task_4), start=1):
            connection.execute(
                text(
                    """
                    UPDATE tasks
                    SET eta_seconds = :eta_seconds,
                        eta_reason = :eta_reason,
                        updated_at = :updated_at
                    WHERE task_id = :task_id
                    """
                ),
                {
                    "task_id": task_id,
                    "eta_seconds": estimate_queue_seconds(position=position, agent_count=2, p50_seconds=60),
                    "eta_reason": eta_reason,
                    "updated_at": datetime.now(UTC).isoformat(),
                },
            )
    engine.dispose()

    runner = FakeCollectorActivityRunner(database_url, tmp_path / "collector", fail_dongchedi_once=True)
    result = asyncio.run(run_single_vehicle_task(task_1, runner))

    assert result == {"task_id": task_1, "status": "completed"}
    detail_response = client.get(f"/api/tasks/{task_1}")
    assert detail_response.status_code == 200
    detail = detail_response.json()
    assert detail["eta_seconds"] is not None
    assert detail["eta_reason"]
    event_types = [event["event_type"] for event in detail["events"]]
    assert "degraded_published" in event_types
    assert "upgraded_to_full" in event_types
    assert event_types.index("degraded_published") < event_types.index("upgraded_to_full")
