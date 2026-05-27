from __future__ import annotations

import os
from pathlib import Path
import sys

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from worker_app.task_reports import generate_task_report_outputs


def _write_input_workbook(path: Path, *, platform: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "raw"
    if platform == "autohome":
        sheet.append(["用户名", "来源链接", "购车地", "发表日期", "最满意", "最不满意", "评价详情"])
        sheet.append(["张三", "https://example.invalid/a", "上海", "2026-05-01", "空间大", "车机卡顿", "空间大，车机卡顿"])
    else:
        sheet.append(["用户名", "来源链接", "购车城市", "发布时间", "评价全文"])
        sheet.append(["李四", "https://example.invalid/b", "北京", "2026-05-02", "动力顺，空间够用"])
    workbook.save(path)


def _write_fake_summary_script(path: Path) -> None:
    path.write_text(
        """
import argparse
from openpyxl import Workbook
parser = argparse.ArgumentParser()
parser.add_argument('--input')
parser.add_argument('--autohome-input')
parser.add_argument('--dcd-input')
parser.add_argument('--output', required=True)
parser.add_argument('--model-name', required=True)
parser.add_argument('--progress-file')
args = parser.parse_args()
wb = Workbook()
ws = wb.active
ws.title = '总览摘要'
ws.append(['模块', '内容'])
ws.append(['平台样本', '汽车之家 1 条 / 懂车帝 1 条'])
for name in ['跨平台对比', '综合业务摘要', '产品机会点']:
    s = wb.create_sheet(name)
    s.append(['维度', '内容'])
    s.append(['核心', '空间好评'])
one = wb.create_sheet('一页纸总结')
one.append(['双平台口碑一页纸总结'])
one.append(['空间好评突出。'])
wb.save(args.output)
""",
        encoding="utf-8",
    )


def _write_fake_wordcloud_script(path: Path) -> None:
    path.write_text(
        """
import argparse
from pathlib import Path
from openpyxl import Workbook
parser = argparse.ArgumentParser()
parser.add_argument('--input')
parser.add_argument('--output-dir', required=True)
parser.add_argument('--model-name', required=True)
parser.add_argument('--font-path')
parser.add_argument('--json', action='store_true')
args = parser.parse_args()
output = Path(args.output_dir)
output.mkdir(parents=True, exist_ok=True)
wb = Workbook()
ws = wb.active
ws.title = 'positive_terms'
ws.append(['term', 'weight'])
ws.append(['空间', 2])
neg = wb.create_sheet('negative_terms')
neg.append(['term', 'weight'])
neg.append(['车机', 1])
wb.save(output / f'{args.model_name}_词云词项清单.xlsx')
(output / f'{args.model_name}_优点词云.png').write_bytes(b'png')
(output / f'{args.model_name}_槽点词云.png').write_bytes(b'png')
print('ok')
""",
        encoding="utf-8",
    )


def test_generate_task_report_outputs_requires_llm_for_full_report(tmp_path: Path) -> None:
    autohome = tmp_path / "raw" / "autohome.xlsx"
    dcd = tmp_path / "raw" / "dcd.xlsx"
    _write_input_workbook(autohome, platform="autohome")
    _write_input_workbook(dcd, platform="dongchedi")

    result = generate_task_report_outputs(
        task_id="task_1",
        model_name="测试车",
        autohome_raw_path=autohome,
        dcd_raw_path=dcd,
        output_root=tmp_path / "outputs",
        summary_script=tmp_path / "missing-summary.py",
        wordcloud_script=tmp_path / "missing-wordcloud.py",
        allow_degraded=False,
        env={**os.environ, "LLM_API_KEY": "", "DEEPSEEK_API_KEY": ""},
    )

    assert result["skipped"] is True
    assert result["failure_category"] == "llm_unavailable"
    assert not (tmp_path / "outputs" / "ai" / "final_report.json").exists()


def test_generate_task_report_outputs_allows_rule_degraded_and_writes_pdf(tmp_path: Path) -> None:
    autohome = tmp_path / "raw" / "autohome.xlsx"
    dcd = tmp_path / "raw" / "dcd.xlsx"
    summary_script = tmp_path / "summary.py"
    wordcloud_script = tmp_path / "wordcloud.py"
    _write_input_workbook(autohome, platform="autohome")
    _write_input_workbook(dcd, platform="dongchedi")
    _write_fake_summary_script(summary_script)
    _write_fake_wordcloud_script(wordcloud_script)

    result = generate_task_report_outputs(
        task_id="task_1",
        model_name="测试车",
        autohome_raw_path=autohome,
        dcd_raw_path=dcd,
        output_root=tmp_path / "outputs",
        summary_script=summary_script,
        wordcloud_script=wordcloud_script,
        allow_degraded=True,
        env={**os.environ, "LLM_API_KEY": "", "DEEPSEEK_API_KEY": "", "HERMES_LLM_MODE": "cli"},
    )

    assert result["skipped"] is False
    assert result["degraded"] is True
    assert Path(result["final_report_path"]).exists()
    assert Path(result["analysis_facts_path"]).exists()
    assert Path(result["pdf_path"]).exists()
    assert Path(result["pdf_path"]).read_bytes().startswith(b"%PDF")
    assert str(Path(result["pdf_path"])) in result["artifact_paths"]
