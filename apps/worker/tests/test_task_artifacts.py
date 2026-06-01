from __future__ import annotations

from datetime import UTC, datetime, timedelta
import os
from pathlib import Path
import subprocess
import sys
import zipfile

from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import create_engine, select, text
from sqlalchemy.orm import sessionmaker

WORKER_ROOT = Path(__file__).resolve().parents[1]
APPS_ROOT = Path(__file__).resolve().parents[2]
API_ROOT = APPS_ROOT / "api"
for root in (WORKER_ROOT, API_ROOT):
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))

from app.models import Base, Task, TaskArtifact
from worker_app.corpus import export_vehicle_merged_raw_workbook, load_platform_state, upsert_platform_rows
from worker_app.task_artifacts import (
    TaskArtifactRecord,
    create_comparison_downloads,
    create_single_task_downloads,
    record_task_artifacts,
)


def _seed_task_session(tmp_path: Path):
    engine = create_engine(f"sqlite+pysqlite:///{tmp_path / 'tasks.db'}", future=True)
    Base.metadata.create_all(engine)
    session_local = sessionmaker(bind=engine, future=True)
    session = session_local()
    session.add(
        Task(
            task_id="task_1",
            task_type="single",
            display_name="测试车",
            status="completed",
            current_stage="completed",
            view_token_hash="view-hash",
            manage_token_hash="manage-hash",
            manage_token_expires_at=datetime.now(UTC) + timedelta(days=7),
        )
    )
    session.commit()
    return session


def _write_workbook(path: Path, sheet_name: str = "raw") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(["车型", "评价"])
    sheet.append(["测试车", path.stem])
    workbook.save(path)


def _seed_worker_task_artifact_session(tmp_path: Path):
    engine = create_engine(f"sqlite+pysqlite:///{tmp_path / 'worker-tasks.db'}", future=True)
    with engine.begin() as connection:
        connection.execute(
            text(
                """
                CREATE TABLE task_artifacts (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    task_id TEXT NOT NULL,
                    artifact_type TEXT NOT NULL,
                    path TEXT NOT NULL,
                    downloadable INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL
                )
                """
            )
        )
        connection.execute(
            text(
                """
                INSERT INTO task_artifacts (task_id, artifact_type, path, downloadable, created_at)
                VALUES ('other_task', 'business_zip', '/tmp/other.zip', 1, '2026-05-26T00:00:00+00:00')
                """
            )
        )
    session_local = sessionmaker(bind=engine, future=True)
    return session_local()


def test_export_vehicle_merged_raw_workbook_writes_two_platform_sheets(tmp_path: Path) -> None:
    database_url = f"sqlite+pysqlite:///{tmp_path / 'corpus.db'}"
    upsert_platform_rows(
        database_url=database_url,
        query="测试车",
        model_name="测试车",
        platform="autohome",
        series_id="8089",
        job_id="job_1",
        rows=[
            {"用户名": "车主A", "发表日期": "2026-05-01", "评价详情": "汽车之家评论1", "来源链接": "https://a.example/1"},
            {"用户名": "车主B", "发表日期": "2026-05-02", "评价详情": "汽车之家评论2", "来源链接": "https://a.example/2"},
        ],
    )
    upsert_platform_rows(
        database_url=database_url,
        query="测试车",
        model_name="测试车",
        platform="dongchedi",
        series_id="25398",
        job_id="job_1",
        rows=[
            {"用户名": "车主C", "发布时间": "2026-05-03", "评价全文": "懂车帝评论1", "来源链接": "https://d.example/1"},
            {"用户名": "车主D", "发布时间": "2026-05-04", "评价全文": "懂车帝评论2", "来源链接": "https://d.example/2"},
            {"用户名": "车主E", "发布时间": "2026-05-05", "评价全文": "懂车帝评论3", "来源链接": "https://d.example/3"},
        ],
    )

    output = tmp_path / "测试车_merged_raw.xlsx"
    counts = export_vehicle_merged_raw_workbook(
        database_url=database_url,
        query="测试车",
        autohome_series_id="8089",
        dongchedi_series_id="25398",
        output_path=output,
    )

    assert counts == {"autohome": 2, "dongchedi": 3}
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["汽车之家", "懂车帝"]
    assert workbook["汽车之家"].max_row == 3
    assert workbook["懂车帝"].max_row == 4
    assert workbook["汽车之家"]["B2"].value == "车主A"
    assert workbook["懂车帝"]["A4"].value == "车主E"


def test_corpus_history_is_shared_by_series_id_instead_of_query_text(tmp_path: Path) -> None:
    database_url = f"sqlite+pysqlite:///{tmp_path / 'corpus.db'}"
    upsert_platform_rows(
        database_url=database_url,
        query="风云T11",
        model_name="风云T11",
        platform="autohome",
        series_id="8089",
        job_id="job_1",
        rows=[
            {
                "用户名": "车主A",
                "发表日期": "2026-05-01",
                "评价详情": "同一条评论",
                "来源链接": "https://k.autohome.com.cn/detail/view_01abc.html?pvareaid=1#fragment",
            }
        ],
    )
    upsert_platform_rows(
        database_url=database_url,
        query="风云 T11",
        model_name="风云 T11",
        platform="autohome",
        series_id="8089",
        job_id="job_2",
        rows=[
            {
                "用户名": "车主A",
                "发表日期": "2026-05-01",
                "评价详情": "同一条评论更新",
                "来源链接": "https://k.autohome.com.cn/detail/view_01abc.html",
            }
        ],
    )

    state = load_platform_state(database_url, query="任意别名", platform="autohome", series_id="8089")
    assert state.existing_count == 1
    assert state.known_links == {"https://k.autohome.com.cn/detail/view_01abc.html"}

    output = tmp_path / "merged.xlsx"
    counts = export_vehicle_merged_raw_workbook(
        database_url=database_url,
        query="另一个查询词",
        autohome_series_id="8089",
        dongchedi_series_id="25398",
        output_path=output,
        platforms=["autohome"],
    )

    workbook = load_workbook(output)
    assert counts == {"autohome": 1}
    assert workbook.sheetnames == ["汽车之家"]
    assert workbook["汽车之家"]["O2"].value == "同一条评论更新"


def test_single_task_downloads_create_zip_and_downloadable_artifact_rows(tmp_path: Path) -> None:
    merged_raw = tmp_path / "测试车_merged_raw.xlsx"
    business_report = tmp_path / "business" / "summary.xlsx"
    business_json = tmp_path / "business" / "final_report.json"
    _write_workbook(merged_raw, "汽车之家")
    business_report.parent.mkdir(parents=True, exist_ok=True)
    business_report.write_text("summary", encoding="utf-8")
    business_json.write_text('{"ok": true}', encoding="utf-8")

    records = create_single_task_downloads(
        task_id="task_1",
        output_dir=tmp_path / "downloads",
        merged_raw_path=merged_raw,
        business_files=[business_report, business_json],
    )

    assert [record.artifact_type for record in records] == ["business_zip", "merged_raw_excel"]
    business_zip = Path(records[0].path)
    assert business_zip.exists()
    with zipfile.ZipFile(business_zip) as archive:
        assert sorted(archive.namelist()) == ["final_report.json", "summary.xlsx"]
    assert Path(records[1].path) == merged_raw

    session = _seed_task_session(tmp_path)
    try:
        record_task_artifacts(session, records)
        artifact_rows = session.execute(select(TaskArtifact).order_by(TaskArtifact.id.asc())).scalars().all()
        assert [(row.artifact_type, row.path, row.downloadable) for row in artifact_rows] == [
            ("business_zip", str(business_zip), True),
            ("merged_raw_excel", str(merged_raw), True),
        ]
    finally:
        session.close()


def test_single_task_downloads_can_write_structured_bundle_and_one_pager_artifact(tmp_path: Path) -> None:
    merged_raw = tmp_path / "测试车_merged_raw.xlsx"
    autohome_raw = tmp_path / "raw" / "ZJ测试车原始口碑.xlsx"
    autohome_validation = tmp_path / "raw" / "ZJ测试车原始口碑.validation.json"
    dcd_raw = tmp_path / "raw" / "DCD口碑_测试车.xlsx"
    dcd_failed_pages = tmp_path / "raw" / "DCD口碑_测试车.failed-pages.json"
    final_report = tmp_path / "ai" / "final_report.json"
    one_pager = tmp_path / "downloads" / "测试车_one_pager.xlsx"
    summary = tmp_path / "summary" / "测试车_双平台口碑摘要.xlsx"
    for workbook_path in (merged_raw, autohome_raw, dcd_raw, one_pager, summary):
        _write_workbook(workbook_path)
    autohome_validation.write_text('{"ok": true}', encoding="utf-8")
    dcd_failed_pages.write_text("[]", encoding="utf-8")
    final_report.parent.mkdir(parents=True, exist_ok=True)
    final_report.write_text('{"headline": "测试车"}', encoding="utf-8")

    records = create_single_task_downloads(
        task_id="task_1",
        output_dir=tmp_path / "downloads",
        merged_raw_path=merged_raw,
        business_files=[summary, final_report],
        one_pager_path=one_pager,
        bundle_entries=[
            (autohome_raw, "autohome/raw.xlsx"),
            (autohome_validation, "autohome/validation.json"),
            (dcd_raw, "dongchedi/raw.xlsx"),
            (dcd_failed_pages, "dongchedi/failed-pages.json"),
            (final_report, "report/final_report.json"),
            (summary, "report/summary.xlsx"),
            (one_pager, "report/one_pager.xlsx"),
            (merged_raw, "merged_raw.xlsx"),
        ],
    )

    assert [record.artifact_type for record in records] == ["business_zip", "merged_raw_excel", "one_pager_excel"]
    with zipfile.ZipFile(records[0].path) as archive:
        assert sorted(archive.namelist()) == [
            "autohome/raw.xlsx",
            "autohome/validation.json",
            "dongchedi/failed-pages.json",
            "dongchedi/raw.xlsx",
            "merged_raw.xlsx",
            "report/final_report.json",
            "report/one_pager.xlsx",
            "report/summary.xlsx",
        ]


def test_comparison_downloads_zip_raw_workbooks_and_business_files_as_downloadable_rows(tmp_path: Path) -> None:
    first_raw = tmp_path / "vehicle_a" / "raw.xlsx"
    second_raw = tmp_path / "vehicle_b" / "raw.xlsx"
    business_file = tmp_path / "business" / "comparison_summary.xlsx"
    _write_workbook(first_raw)
    _write_workbook(second_raw)
    _write_workbook(business_file)

    records = create_comparison_downloads(
        task_id="task_1",
        output_dir=tmp_path / "downloads",
        vehicle_raw_paths={"车辆A": first_raw, "车辆B": second_raw},
        business_files=[business_file],
    )

    assert [record.artifact_type for record in records] == ["business_zip", "vehicle_raw_excel", "vehicle_raw_excel"]
    comparison_zip = Path(records[0].path)
    with zipfile.ZipFile(comparison_zip) as archive:
        assert sorted(archive.namelist()) == [
            "comparison_summary.xlsx",
            "vehicle_raw/车辆A_raw.xlsx",
            "vehicle_raw/车辆B_raw.xlsx",
        ]
    assert [Path(record.path) for record in records[1:]] == [first_raw, second_raw]

    session = _seed_task_session(tmp_path)
    try:
        record_task_artifacts(session, records)
        rows = session.execute(select(TaskArtifact).order_by(TaskArtifact.id.asc())).scalars().all()
        assert [row.downloadable for row in rows] == [True, True, True]
        assert [row.artifact_type for row in rows] == ["business_zip", "vehicle_raw_excel", "vehicle_raw_excel"]
    finally:
        session.close()


def test_record_task_artifacts_is_idempotent_with_worker_schema_without_app_models(tmp_path: Path) -> None:
    session = _seed_worker_task_artifact_session(tmp_path)
    records = [
        TaskArtifactRecord("task_1", "business_zip", "/tmp/task_1.zip"),
        TaskArtifactRecord("task_1", "merged_raw_excel", "/tmp/task_1_raw.xlsx"),
    ]

    try:
        record_task_artifacts(session, records)
        record_task_artifacts(session, records)
        rows = session.execute(
            text("SELECT task_id, artifact_type, path, downloadable FROM task_artifacts ORDER BY task_id, artifact_type")
        ).all()
    finally:
        session.close()

    assert rows == [
        ("other_task", "business_zip", "/tmp/other.zip", 1),
        ("task_1", "business_zip", "/tmp/task_1.zip", 1),
        ("task_1", "merged_raw_excel", "/tmp/task_1_raw.xlsx", 1),
    ]


def test_task_artifacts_module_imports_and_records_with_worker_pythonpath_only(tmp_path: Path) -> None:
    worker_root = Path(__file__).resolve().parents[1]
    db_path = tmp_path / "worker-only.db"
    script = f"""
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker
from worker_app.task_artifacts import TaskArtifactRecord, record_task_artifacts

engine = create_engine("sqlite+pysqlite:///{db_path}", future=True)
with engine.begin() as connection:
    connection.execute(text('''
        CREATE TABLE task_artifacts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            task_id TEXT NOT NULL,
            artifact_type TEXT NOT NULL,
            path TEXT NOT NULL,
            downloadable INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL
        )
    '''))
session = sessionmaker(bind=engine, future=True)()
record_task_artifacts(session, [TaskArtifactRecord("task_1", "business_zip", "/tmp/task.zip")])
session.close()
"""

    env = {**os.environ, "PYTHONPATH": str(worker_root)}
    result = subprocess.run(
        [sys.executable, "-c", script],
        cwd=worker_root,
        env=env,
        check=False,
        capture_output=True,
        text=True,
    )

    assert result.returncode == 0, result.stderr
    engine = create_engine(f"sqlite+pysqlite:///{db_path}", future=True)
    with engine.begin() as connection:
        rows = connection.execute(text("SELECT task_id, artifact_type, path, downloadable FROM task_artifacts")).all()
    assert rows == [("task_1", "business_zip", "/tmp/task.zip", 1)]


def test_single_task_downloads_reject_missing_files_without_partial_zip(tmp_path: Path) -> None:
    merged_raw = tmp_path / "missing_raw.xlsx"
    business_file = tmp_path / "business" / "summary.xlsx"
    business_file.parent.mkdir(parents=True, exist_ok=True)
    business_file.write_text("summary", encoding="utf-8")

    with pytest.raises(FileNotFoundError, match=str(merged_raw)):
        create_single_task_downloads("task_1", tmp_path / "downloads", merged_raw, [business_file])
    assert not (tmp_path / "downloads" / "task_1_business.zip").exists()

    _write_workbook(merged_raw)
    missing_business = tmp_path / "business" / "missing.xlsx"
    with pytest.raises(FileNotFoundError, match=str(missing_business)):
        create_single_task_downloads("task_1", tmp_path / "downloads", merged_raw, [missing_business])
    assert not (tmp_path / "downloads" / "task_1_business.zip").exists()


def test_comparison_downloads_reject_missing_files_without_partial_zip(tmp_path: Path) -> None:
    first_raw = tmp_path / "vehicle_a" / "raw.xlsx"
    missing_raw = tmp_path / "vehicle_b" / "raw.xlsx"
    business_file = tmp_path / "business" / "comparison_summary.xlsx"
    _write_workbook(first_raw)
    _write_workbook(business_file)

    with pytest.raises(FileNotFoundError, match=str(missing_raw)):
        create_comparison_downloads(
            task_id="task_1",
            output_dir=tmp_path / "downloads",
            vehicle_raw_paths={"车辆A": first_raw, "车辆B": missing_raw},
            business_files=[business_file],
        )
    assert not (tmp_path / "downloads" / "task_1_business.zip").exists()

    missing_business = tmp_path / "business" / "missing.xlsx"
    with pytest.raises(FileNotFoundError, match=str(missing_business)):
        create_comparison_downloads(
            task_id="task_1",
            output_dir=tmp_path / "downloads",
            vehicle_raw_paths={"车辆A": first_raw},
            business_files=[missing_business],
        )
    assert not (tmp_path / "downloads" / "task_1_business.zip").exists()
